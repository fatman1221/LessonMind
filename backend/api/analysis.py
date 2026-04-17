from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from pydantic import BaseModel
from typing import List, Dict, Any, Optional
from datetime import datetime, timedelta
import os
from urllib.parse import quote

from database import get_db
from models import (
    User,
    StudentData,
    QuestionBank,
    AssignmentSubmission,
    Assignment,
    Class,
    ClassStudent,
    ChatHistory,
    AlertPreference,
)
from services.ai_service import ai_service
from services.analysis_service import analysis_service
from api.auth import get_current_user
from api.assignment import grade_answers, _normalize_assignment_questions
from config import settings

router = APIRouter()


class StudentDataCreate(BaseModel):
    student_id: str
    student_name: str
    subject: str
    grade: str
    homework_scores: Dict[str, float]
    learning_behavior: Dict[str, Any]
    knowledge_mastery: Dict[str, float]


class AnalysisResult(BaseModel):
    mastery_level: str
    mastery_score: float
    ability_score: float
    habit_score: float
    prediction_confidence: float
    recommendations: List[str]


class QuestionGenerateRequest(BaseModel):
    knowledge_point: str
    subject: str
    question_types: List[str] = ["choice", "fill", "short_answer"]
    count: int = 3


class QuestionItem(BaseModel):
    type: str
    question: str
    options: Optional[List[str]] = None
    answer: str
    explanation: str


class QuestionBankItem(BaseModel):
    id: int
    subject: str
    knowledge_point: str
    question_type: str
    question_content: str
    answer: str
    difficulty: str
    created_at: datetime
    
    class Config:
        from_attributes = True


class ClassAssignmentOverviewItem(BaseModel):
    class_id: int
    class_name: str
    student_count: int
    assignment_count: int
    submission_count: int
    completion_rate: float
    avg_score_pct: float


class TeachingReportResponse(BaseModel):
    teacher_id: int
    classes: List[ClassAssignmentOverviewItem]
    student_score_overview: Dict[str, float]
    usage_frequency: Dict[str, Any]
    preference_tags: List[str]
    prep_recommendations: List[str]


class StudentRecommendationItem(BaseModel):
    question_id: int
    question_type: str
    question_content: str
    reason: str


class StudentVideoRecommendationItem(BaseModel):
    title: str
    url: str
    reason: str


class StudentRecommendationResponse(BaseModel):
    student_id: int
    signals: Dict[str, Any]
    recommended_questions: List[StudentRecommendationItem]
    recommended_videos: List[StudentVideoRecommendationItem]


class TeachingAlertItem(BaseModel):
    level: str  # high / medium / low
    category: str  # deadline / class_completion / student_risk
    title: str
    detail: str
    action_hint: str
    alert_key: str
    class_id: Optional[int] = None
    assignment_id: Optional[int] = None
    student_id: Optional[int] = None


class AlertStateUpdateRequest(BaseModel):
    alert_key: str
    ignored: bool = True


@router.post("/student-data")
async def create_student_data(
    student_data: StudentDataCreate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """导入学生数据"""
    try:
        db_student = StudentData(
            student_id=student_data.student_id,
            student_name=student_data.student_name,
            subject=student_data.subject,
            grade=student_data.grade,
            homework_scores=student_data.homework_scores,
            learning_behavior=student_data.learning_behavior,
            knowledge_mastery=student_data.knowledge_mastery
        )
        db.add(db_student)
        db.commit()
        db.refresh(db_student)
        
        return {"message": "学生数据导入成功", "id": db_student.id}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"导入学生数据失败: {str(e)}")


@router.post("/analyze/{student_id}", response_model=AnalysisResult)
async def analyze_student(
    student_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """分析学生学情（优先使用导入的 StudentData；否则根据作业提交自动生成）"""
    try:
        sid = int(student_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="无效的学生ID")

    student = db.query(User).filter(User.id == sid).first()
    if not student:
        raise HTTPException(status_code=404, detail="学生用户不存在")
    if student.role != "student":
        raise HTTPException(status_code=400, detail="仅支持分析学生账号")

    student_data = db.query(StudentData).filter(StudentData.student_id == student_id).first()

    if student_data:
        analysis_data = {
            "homework_scores": student_data.homework_scores or {},
            "learning_behavior": student_data.learning_behavior or {},
            "knowledge_mastery": student_data.knowledge_mastery or {},
        }
        from_import = True
    else:
        analysis_data = _analysis_data_from_submissions(db, sid)
        from_import = False

    result = analysis_service.analyze_student(analysis_data)
    signals = _student_dynamic_signals(db, sid)
    recs = list(result.get("recommendations") or [])

    if not from_import and not analysis_data.get("homework_scores"):
        extra = "该学生暂无作业记录，建议自动推送诊断测试与对应教学视频，并在24小时后复测。"
        if extra not in recs:
            recs.insert(0, extra)
    elif signals["completion_rate"] < 0.6:
        recs.append("作业完成率偏低，建议安排分层作业并设置阶段提醒。")

    if signals["chat_frequency_7d"] < 2:
        recs.append("近7天学习互动频次较低，建议在课堂中增加问答与同伴讨论环节。")
    if signals["covered_topics"] < 3:
        recs.append("章节学习覆盖面较窄，建议补充跨章节综合练习与复盘。")

    result = {**result, "recommendations": list(dict.fromkeys(recs))}

    return AnalysisResult(**result)


@router.post("/train-model")
async def train_analysis_model(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """训练学情分析模型"""
    try:
        # 获取所有学生数据
        all_students = db.query(StudentData).all()
        
        if len(all_students) < 10:
            raise HTTPException(
                status_code=400,
                detail="训练数据不足，至少需要10条学生数据"
            )
        
        # 准备训练数据
        training_data = []
        for student in all_students:
            training_data.append({
                "homework_scores": student.homework_scores,
                "learning_behavior": student.learning_behavior,
                "knowledge_mastery": student.knowledge_mastery
            })
        
        # 训练模型
        analysis_service.train_model(training_data)
        
        return {"message": "模型训练成功", "training_samples": len(training_data)}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"模型训练失败: {str(e)}")


@router.get("/class-overview", response_model=List[ClassAssignmentOverviewItem])
async def get_class_assignment_overview(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """班级与作业关联概览，便于教师分析。"""
    if current_user.role not in ("teacher", "admin"):
        raise HTTPException(status_code=403, detail="仅教师或管理员可访问")

    class_query = db.query(Class)
    if current_user.role == "teacher":
        class_query = class_query.filter(Class.teacher_id == current_user.id)
    classes = class_query.order_by(Class.created_at.desc()).all()

    result: List[ClassAssignmentOverviewItem] = []
    for cls in classes:
        student_count = db.query(ClassStudent).filter(ClassStudent.class_id == cls.id).count()
        assignments = db.query(Assignment).filter(Assignment.class_id == cls.id).all()
        assignment_ids = [a.id for a in assignments]
        assignment_count = len(assignment_ids)
        if assignment_ids:
            submissions = db.query(AssignmentSubmission).filter(AssignmentSubmission.assignment_id.in_(assignment_ids)).all()
        else:
            submissions = []
        submission_count = len(submissions)

        expected = max(student_count * assignment_count, 1)
        completion_rate = submission_count / expected if expected else 0.0
        score_pcts: List[float] = []
        for s in submissions:
            if s.total_score and float(s.total_score) > 0:
                score_pcts.append(float(s.score or 0) / float(s.total_score))
        avg_score_pct = float(sum(score_pcts) / len(score_pcts) * 100.0) if score_pcts else 0.0

        result.append(ClassAssignmentOverviewItem(
            class_id=cls.id,
            class_name=cls.name,
            student_count=student_count,
            assignment_count=assignment_count,
            submission_count=submission_count,
            completion_rate=round(completion_rate, 4),
            avg_score_pct=round(avg_score_pct, 2),
        ))

    return result


@router.get("/teaching-report", response_model=TeachingReportResponse)
async def get_teaching_report(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """教师备课报告：成绩、使用频率、偏好与建议。"""
    if current_user.role not in ("teacher", "admin"):
        raise HTTPException(status_code=403, detail="仅教师或管理员可访问")

    teacher_id = current_user.id
    report = _build_teacher_report(db, teacher_id)
    return TeachingReportResponse(**report)


@router.get("/student-recommendations/{student_id}", response_model=StudentRecommendationResponse)
async def get_student_recommendations(
    student_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """根据学生动态监测结果自动推送测试题和教学视频建议。"""
    if current_user.role not in ("teacher", "admin"):
        raise HTTPException(status_code=403, detail="仅教师或管理员可访问")

    student = db.query(User).filter(User.id == student_id).first()
    if not student or student.role != "student":
        raise HTTPException(status_code=404, detail="学生不存在")

    # 教师仅可查看自己班级中的学生
    if current_user.role == "teacher":
        owns = db.query(ClassStudent).join(Class, Class.id == ClassStudent.class_id).filter(
            ClassStudent.student_id == student_id,
            Class.teacher_id == current_user.id
        ).first()
        if not owns:
            raise HTTPException(status_code=403, detail="无权查看该学生")

    signals = _student_dynamic_signals(db, student_id)
    profile = _infer_student_profile(db, student_id)
    subject = profile.get("subject")

    query = db.query(QuestionBank)
    if subject:
        query = query.filter(QuestionBank.subject == subject)
    question_rows = query.order_by(QuestionBank.created_at.desc()).limit(6).all()

    reason = "用于巩固薄弱知识点与提升作业完成质量"
    if signals.get("completion_rate", 0) < 0.6:
        reason = "该生作业完成率偏低，建议先做短小练习建立节奏"

    recommended_questions = [
        StudentRecommendationItem(
            question_id=q.id,
            question_type=q.question_type,
            question_content=q.question_content[:120],
            reason=reason,
        )
        for q in question_rows
    ]

    video_topics = [
        f"{subject or '通用'} 基础概念精讲",
        f"{subject or '通用'} 典型题讲解",
        f"{subject or '通用'} 易错点复盘",
    ]
    if signals.get("chat_frequency_7d", 0) < 2:
        video_topics.append(f"{subject or '通用'} 高效学习方法")

    recommended_videos = [
        StudentVideoRecommendationItem(
            title=t,
            url=f"https://www.bilibili.com/search?keyword={quote(t)}",
            reason="建议课后观看并完成配套练习",
        )
        for t in video_topics
    ]

    return StudentRecommendationResponse(
        student_id=student_id,
        signals=signals,
        recommended_questions=recommended_questions,
        recommended_videos=recommended_videos,
    )


@router.get("/teaching-alerts", response_model=List[TeachingAlertItem])
async def get_teaching_alerts(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """教师提醒中心：截止提醒、班级完成率预警、重点关注学生。"""
    if current_user.role not in ("teacher", "admin"):
        raise HTTPException(status_code=403, detail="仅教师或管理员可访问")

    if current_user.role == "teacher":
        classes = db.query(Class).filter(Class.teacher_id == current_user.id).all()
        assignments = db.query(Assignment).filter(Assignment.teacher_id == current_user.id).all()
    else:
        classes = db.query(Class).all()
        assignments = db.query(Assignment).all()

    now = datetime.utcnow()
    alerts: List[TeachingAlertItem] = []

    # 1) 即将截止提醒（24小时内）
    for a in assignments:
        if not a.is_published or not a.deadline:
            continue
        delta = (a.deadline - now).total_seconds()
        if 0 < delta <= 24 * 3600:
            alerts.append(TeachingAlertItem(
                level="high",
                category="deadline",
                title=f"作业《{a.title}》即将截止",
                detail=f"距离截止约 {int(delta // 3600)} 小时，建议立即提醒学生完成提交。",
                action_hint="可在班级群发送截止提醒，并开启补交策略。",
                alert_key=f"deadline-{a.id}",
                class_id=a.class_id,
                assignment_id=a.id,
            ))

    # 2) 班级完成率预警
    for cls in classes:
        members = db.query(ClassStudent).filter(ClassStudent.class_id == cls.id).all()
        student_ids = [m.student_id for m in members]
        class_assignments = [a for a in assignments if a.class_id == cls.id]
        assign_ids = [a.id for a in class_assignments]
        submissions = db.query(AssignmentSubmission).filter(
            AssignmentSubmission.assignment_id.in_(assign_ids)
        ).all() if assign_ids else []
        expected = max(len(student_ids) * len(class_assignments), 1)
        completion_rate = (len(submissions) / expected) if expected else 0.0
        if len(class_assignments) > 0 and completion_rate < 0.6:
            alerts.append(TeachingAlertItem(
                level="medium",
                category="class_completion",
                title=f"{cls.name} 作业完成率偏低",
                detail=f"当前完成率约 {round(completion_rate * 100)}%，建议进行分层督导。",
                action_hint="优先跟进未提交学生，分配短任务提升完成度。",
                alert_key=f"class-completion-{cls.id}",
                class_id=cls.id,
            ))

    # 3) 重点关注学生（按完成率）
    class_ids = [c.id for c in classes]
    student_rows = db.query(ClassStudent).filter(ClassStudent.class_id.in_(class_ids)).all() if class_ids else []
    student_ids = sorted({s.student_id for s in student_rows})
    for sid in student_ids[:80]:  # 限制范围，避免开销过大
        signals = _student_dynamic_signals(db, sid)
        if signals.get("completion_rate", 0) < 0.5:
            stu = db.query(User).filter(User.id == sid).first()
            if not stu:
                continue
            alerts.append(TeachingAlertItem(
                level="medium",
                category="student_risk",
                title=f"学生 {stu.username} 需要重点关注",
                detail=f"作业完成率 {round(signals.get('completion_rate', 0) * 100)}%，近7天互动 {signals.get('chat_frequency_7d', 0)} 次。",
                action_hint="建议推送诊断题+视频，并安排一次课后沟通。",
                alert_key=f"student-risk-{sid}",
                student_id=sid,
            ))

    # 控制返回数量，按严重级别排序
    level_order = {"high": 0, "medium": 1, "low": 2}
    alerts.sort(key=lambda x: level_order.get(x.level, 9))
    return alerts[:20]


@router.get("/teaching-alerts/ignored", response_model=List[str])
async def get_ignored_teaching_alerts(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """获取当前用户已忽略的提醒键。"""
    if current_user.role not in ("teacher", "admin"):
        raise HTTPException(status_code=403, detail="仅教师或管理员可访问")

    rows = db.query(AlertPreference).filter(
        AlertPreference.user_id == current_user.id,
        AlertPreference.is_ignored == True
    ).all()
    return [r.alert_key for r in rows]


@router.post("/teaching-alerts/state")
async def update_teaching_alert_state(
    req: AlertStateUpdateRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """更新提醒忽略状态。"""
    if current_user.role not in ("teacher", "admin"):
        raise HTTPException(status_code=403, detail="仅教师或管理员可访问")

    row = db.query(AlertPreference).filter(
        AlertPreference.user_id == current_user.id,
        AlertPreference.alert_key == req.alert_key
    ).first()

    if row:
        row.is_ignored = bool(req.ignored)
    else:
        row = AlertPreference(
            user_id=current_user.id,
            alert_key=req.alert_key,
            is_ignored=bool(req.ignored),
        )
        db.add(row)

    db.commit()
    return {"message": "状态已更新", "alert_key": req.alert_key, "ignored": bool(req.ignored)}


@router.delete("/teaching-alerts/ignored")
async def clear_ignored_teaching_alerts(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """清空当前用户所有忽略提醒记录。"""
    if current_user.role not in ("teacher", "admin"):
        raise HTTPException(status_code=403, detail="仅教师或管理员可访问")

    rows = db.query(AlertPreference).filter(AlertPreference.user_id == current_user.id).all()
    count = 0
    for row in rows:
        if row.is_ignored:
            row.is_ignored = False
            count += 1
    db.commit()
    return {"message": f"已清空 {count} 条忽略记录"}


@router.post("/questions/generate", response_model=List[QuestionItem])
async def generate_questions(
    request: QuestionGenerateRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """生成题目"""
    try:
        # 调用AI服务生成题目
        questions = ai_service.generate_questions(
            knowledge_point=request.knowledge_point,
            subject=request.subject,
            question_types=request.question_types,
            count=request.count
        )
        
        # 保存到题库
        for q in questions:
            question_bank = QuestionBank(
                subject=request.subject,
                knowledge_point=request.knowledge_point,
                question_type=q.get("type", "choice"),
                question_content=q.get("question", ""),
                answer=q.get("answer", ""),
                difficulty="medium"
            )
            db.add(question_bank)
        
        db.commit()
        
        return questions
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"生成题目失败: {str(e)}")


@router.get("/questions", response_model=List[QuestionBankItem])
async def get_questions(
    subject: Optional[str] = None,
    knowledge_point: Optional[str] = None,
    question_type: Optional[str] = None,
    skip: int = 0,
    limit: int = 50,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """获取题库题目"""
    query = db.query(QuestionBank)
    
    if subject:
        query = query.filter(QuestionBank.subject == subject)
    if knowledge_point:
        query = query.filter(QuestionBank.knowledge_point.contains(knowledge_point))
    if question_type:
        query = query.filter(QuestionBank.question_type == question_type)
    
    questions = query.order_by(QuestionBank.created_at.desc())\
        .offset(skip)\
        .limit(limit)\
        .all()
    
    return questions


@router.post("/questions/export")
async def export_questions(
    question_ids: List[int],
    format: str = "json",
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """导出题目"""
    questions = db.query(QuestionBank)\
        .filter(QuestionBank.id.in_(question_ids))\
        .all()
    
    if not questions:
        raise HTTPException(status_code=404, detail="未找到题目")
    
    if format == "json":
        result = []
        for q in questions:
            result.append({
                "id": q.id,
                "subject": q.subject,
                "knowledge_point": q.knowledge_point,
                "question_type": q.question_type,
                "question_content": q.question_content,
                "answer": q.answer,
                "difficulty": q.difficulty
            })
        return {"format": "json", "data": result}
    else:
        # 可以扩展为其他格式（如Word、Excel）
        raise HTTPException(status_code=400, detail="不支持的导出格式")


@router.post("/questions/export-word")
async def export_questions_to_word(
    question_ids: List[int],
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """导出题目为Word文档"""
    from fastapi.responses import FileResponse
    from services.multimedia_service import multimedia_service
    import os
    from datetime import datetime
    
    questions = db.query(QuestionBank)\
        .filter(QuestionBank.id.in_(question_ids))\
        .all()
    
    if not questions:
        raise HTTPException(status_code=404, detail="未找到题目")
    
    # 准备题目数据
    questions_data = []
    for q in questions:
        questions_data.append({
            "subject": q.subject,
            "knowledge_point": q.knowledge_point,
            "question_type": q.question_type,
            "question_content": q.question_content,
            "answer": q.answer,
            "difficulty": q.difficulty
        })
    
    # 生成Word文档
    word_dir = os.path.join(settings.UPLOAD_DIR, "word")
    os.makedirs(word_dir, exist_ok=True)
    
    filename = f"question_bank_{datetime.now().strftime('%Y%m%d%H%M%S')}.docx"
    file_path = os.path.join(word_dir, filename)
    
    # 创建Word文档
    from docx import Document
    from docx.shared import Pt, RGBColor as DocxRGBColor
    
    doc = Document()
    
    # 设置文档样式
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(12)
    
    # 添加标题
    title = doc.add_heading('题库导出', 0)
    try:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except:
        pass
    
    # 添加题目
    for i, q in enumerate(questions_data, 1):
        # 题目编号和类型
        doc.add_heading(f'第{i}题 ({get_question_type_name(q["question_type"])})', 2)
        
        # 题目内容
        p = doc.add_paragraph()
        p.add_run('题目：').bold = True
        p.add_run(q['question_content'])
        
        # 答案
        p = doc.add_paragraph()
        p.add_run('答案：').bold = True
        p.add_run(q['answer'])
        
        # 分隔线
        if i < len(questions_data):
            doc.add_paragraph('_' * 50)
    
    # 保存文档
    doc.save(file_path)
    
    return FileResponse(
        file_path,
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        filename=filename
    )


@router.delete("/questions/{question_id}")
async def delete_question(
    question_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """删除题目"""
    question = db.query(QuestionBank).filter(QuestionBank.id == question_id).first()
    
    if not question:
        raise HTTPException(status_code=404, detail="题目不存在")
    
    db.delete(question)
    db.commit()
    
    return {"message": "删除成功"}


@router.post("/questions/delete-batch")
async def delete_questions_batch(
    question_ids: List[int],
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """批量删除题目"""
    if not question_ids:
        raise HTTPException(status_code=400, detail="请选择要删除的题目")
    
    questions = db.query(QuestionBank).filter(QuestionBank.id.in_(question_ids)).all()
    
    if not questions:
        raise HTTPException(status_code=404, detail="未找到要删除的题目")
    
    for question in questions:
        db.delete(question)
    
    db.commit()
    
    return {"message": f"成功删除 {len(questions)} 道题目"}


def get_question_type_name(q_type: str) -> str:
    """获取题目类型名称"""
    type_map = {
        "choice": "选择题",
        "fill": "填空题",
        "short_answer": "简答题"
    }
    return type_map.get(q_type, q_type)


def _analysis_data_from_submissions(db: Session, student_db_id: int) -> Dict[str, Any]:
    """
    无「学生数据」导入记录时，根据作业提交记录构造分析输入。
    homework_scores 使用 0–100 分制，与原有 StudentData 导入格式一致。
    """
    subs = (
        db.query(AssignmentSubmission)
        .filter(AssignmentSubmission.student_id == student_db_id)
        .order_by(AssignmentSubmission.submitted_at.desc())
        .all()
    )
    homework_scores: Dict[str, float] = {}
    correct_total = 0
    question_total = 0

    for sub in subs:
        assignment = (
            db.query(Assignment).filter(Assignment.id == sub.assignment_id).first()
        )
        title = (assignment.title if assignment else f"作业{sub.assignment_id}")[:80]
        if title in homework_scores:
            title = f"{title}({sub.assignment_id})"
        if sub.total_score and float(sub.total_score) > 0:
            pct = float(sub.score or 0) / float(sub.total_score) * 100.0
        else:
            pct = 0.0
        homework_scores[title] = round(pct, 1)

        if assignment:
            questions = _normalize_assignment_questions(assignment.questions)
            if questions:
                gr = grade_answers(questions, sub.answers or {})
                correct_total += int(gr.get("correct_count", 0))
                question_total += int(gr.get("total_questions", len(questions)))

    n = len(subs)
    learning_behavior = {
        "study_time": min(n * 2, 20),
        "question_count": max(n * 2, n),
        "participation_rate": round(min(1.0, 0.35 + n * 0.12), 2) if n else 0.25,
    }

    knowledge_mastery: Dict[str, float] = {}
    if question_total > 0:
        knowledge_mastery["作业题目正确率"] = round(correct_total / question_total, 4)
    elif homework_scores:
        knowledge_mastery["作业得分率"] = round(
            sum(homework_scores.values()) / len(homework_scores) / 100.0, 4
        )

    return {
        "homework_scores": homework_scores,
        "learning_behavior": learning_behavior,
        "knowledge_mastery": knowledge_mastery,
    }


def _student_dynamic_signals(db: Session, student_id: int) -> Dict[str, Any]:
    """动态监测：作业完成、章节学习（覆盖面）、对话频次。"""
    class_ids = [c[0] for c in db.query(ClassStudent.class_id).filter(ClassStudent.student_id == student_id).all()]
    assignment_query = db.query(Assignment)
    if class_ids:
        assignment_query = assignment_query.filter(Assignment.class_id.in_(class_ids))
    assignments = assignment_query.all()
    total_assignments = len(assignments)
    assignment_ids = [a.id for a in assignments]

    submitted = db.query(AssignmentSubmission).filter(AssignmentSubmission.student_id == student_id)
    if assignment_ids:
        submitted = submitted.filter(AssignmentSubmission.assignment_id.in_(assignment_ids))
    submitted_rows = submitted.all()
    completion_rate = (len(submitted_rows) / total_assignments) if total_assignments else 0.0

    covered_topics = set()
    for a in assignments:
        qs = _normalize_assignment_questions(a.questions)
        for q in qs:
            text = str(q.get("question", "")).strip()
            if text:
                covered_topics.add(text[:24])

    seven_days_ago = datetime.utcnow() - timedelta(days=7)
    chat_count_7d = db.query(ChatHistory).filter(
        ChatHistory.user_id == student_id,
        ChatHistory.created_at >= seven_days_ago
    ).count()

    return {
        "completion_rate": round(completion_rate, 4),
        "covered_topics": len(covered_topics),
        "chat_frequency_7d": int(chat_count_7d),
    }


def _build_teacher_report(db: Session, teacher_id: int) -> Dict[str, Any]:
    classes = db.query(Class).filter(Class.teacher_id == teacher_id).all()
    class_items: List[Dict[str, Any]] = []
    all_scores: List[float] = []
    all_completion: List[float] = []

    student_ids = set()
    for cls in classes:
        members = db.query(ClassStudent).filter(ClassStudent.class_id == cls.id).all()
        cls_student_ids = [m.student_id for m in members]
        student_ids.update(cls_student_ids)

        assignments = db.query(Assignment).filter(Assignment.class_id == cls.id).all()
        assignment_ids = [a.id for a in assignments]
        submissions = db.query(AssignmentSubmission).filter(AssignmentSubmission.assignment_id.in_(assignment_ids)).all() if assignment_ids else []

        expected = max(len(cls_student_ids) * len(assignments), 1)
        completion_rate = (len(submissions) / expected) if expected else 0.0
        score_pct = [
            (float(s.score or 0) / float(s.total_score)) * 100.0
            for s in submissions if s.total_score and float(s.total_score) > 0
        ]
        avg_score = (sum(score_pct) / len(score_pct)) if score_pct else 0.0

        class_items.append({
            "class_id": cls.id,
            "class_name": cls.name,
            "student_count": len(cls_student_ids),
            "assignment_count": len(assignments),
            "submission_count": len(submissions),
            "completion_rate": round(completion_rate, 4),
            "avg_score_pct": round(avg_score, 2),
        })
        all_completion.append(completion_rate)
        all_scores.extend(score_pct)

    # 使用频率：教师聊天与资源产出
    chat_rows = db.query(ChatHistory).filter(ChatHistory.user_id == teacher_id).all()
    now_ts = datetime.utcnow().timestamp()
    week_ago = now_ts - 7 * 24 * 3600
    chat_7d = sum(1 for c in chat_rows if c.created_at and c.created_at.timestamp() >= week_ago)
    avg_daily_chat = round(chat_7d / 7.0, 2)

    # 自由偏好：从聊天内容提取高频教学关键词
    keywords = ["数学", "语文", "英语", "物理", "化学", "生物", "历史", "地理", "实验", "提问", "分层", "互动", "作业"]
    pref_counter: Dict[str, int] = {}
    for row in chat_rows:
        msg = row.message or ""
        for kw in keywords:
            if kw in msg:
                pref_counter[kw] = pref_counter.get(kw, 0) + 1
    preference_tags = [k for k, _ in sorted(pref_counter.items(), key=lambda x: x[1], reverse=True)[:6]]

    avg_completion = round(sum(all_completion) / len(all_completion), 4) if all_completion else 0.0
    avg_score = round(sum(all_scores) / len(all_scores), 2) if all_scores else 0.0
    prep_recommendations: List[str] = []
    if avg_completion < 0.7:
        prep_recommendations.append("班级作业完成率偏低，建议备课时增加分层任务与阶段提醒。")
    if avg_score < 70:
        prep_recommendations.append("平均成绩偏低，建议在备课中加入基础巩固与错题复盘环节。")
    if avg_daily_chat < 1:
        prep_recommendations.append("教师系统使用频率较低，建议结合AI问答进行教案与题目预演。")
    if not prep_recommendations:
        prep_recommendations.append("当前教学节奏较稳定，可增加跨章节综合任务提升迁移能力。")

    return {
        "teacher_id": teacher_id,
        "classes": class_items,
        "student_score_overview": {
            "avg_score_pct": avg_score,
            "avg_completion_rate": avg_completion,
            "student_count": float(len(student_ids)),
        },
        "usage_frequency": {
            "chat_count_7d": chat_7d,
            "avg_daily_chat_7d": avg_daily_chat,
        },
        "preference_tags": preference_tags,
        "prep_recommendations": prep_recommendations,
    }


def _infer_student_profile(db: Session, student_id: int) -> Dict[str, Optional[str]]:
    """从班级关系推断学生的学科与学段。"""
    cls = db.query(Class).join(ClassStudent, Class.id == ClassStudent.class_id).filter(
        ClassStudent.student_id == student_id
    ).order_by(Class.created_at.desc()).first()
    if not cls:
        return {"subject": None, "grade": None}
    return {"subject": cls.subject, "grade": cls.grade}

